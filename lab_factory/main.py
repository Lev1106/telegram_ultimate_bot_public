import html
import io
import json
import os
import re
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pymupdf
from PIL import Image, UnidentifiedImageError
from dotenv import load_dotenv
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from docx.text.paragraph import Paragraph
from google import genai
from google.genai import types

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image as RLImage,
    PageBreak,
    Paragraph as RLParagraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table as RLTable,
    TableStyle as RLTableStyle,
)


BASE_DIR = Path(__file__).parent
INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".webp"}
ASSIGNMENT_EXTENSIONS = {".txt", ".docx", ".pdf"}
CODE_EXTENSIONS = {
    ".py", ".cpp", ".cc", ".cxx", ".c", ".h", ".hpp",
    ".java", ".kt", ".kts",
    ".js", ".jsx", ".ts", ".tsx",
    ".cs", ".go", ".rs", ".php", ".swift",
    ".sql", ".sh", ".bash", ".zsh",
    ".html", ".htm", ".css", ".scss",
    ".json", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg",
    ".md", ".r", ".m", ".pl", ".lua",
}
SPECIAL_ROOT_FILES = {"meta.json", "notes.txt"}

INLINE_PATTERN = re.compile(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)")
LIST_ITEM_RE = re.compile(r"^\s*(\d+[\.\)]|[-•])\s+(.*)$")
PIPE_TABLE_SEPARATOR_RE = re.compile(
    r"^\s*\|?(?:\s*:?-{3,}:?\s*\|)+\s*:?-{3,}:?\s*\|?\s*$"
)

GEMINI_MODELS_FALLBACK = [
    "gemini-2.5-flash-lite",
    "gemini-2.5-flash",
    "gemini-3.1-flash-lite-preview",
    "gemini-3-flash-preview",
]

MAX_IMAGES_FOR_ANALYSIS = 6
MAX_IMAGE_DIMENSION_FOR_ANALYSIS = 1600


# -----------------------------
# FILE READING
# -----------------------------
def read_text_file(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8").strip()


def read_text_file_safe(path: Path) -> str:
    if not path.exists():
        return ""

    encodings = ["utf-8", "utf-8-sig", "cp1251", "latin-1"]
    for enc in encodings:
        try:
            return path.read_text(encoding=enc).strip()
        except Exception:
            continue

    try:
        return path.read_text(encoding="utf-8", errors="ignore").strip()
    except Exception:
        return ""


def read_json_file(path: Path) -> Dict[str, Any]:
    if not path.exists():
        return {}
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def read_docx_file(path: Path) -> str:
    if not path.exists():
        return ""

    doc = Document(path)
    parts: List[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_values = []
            for cell in row.cells:
                value = cell.text.strip()
                row_values.append(value)
            if any(row_values):
                parts.append(" | ".join(row_values))

    return "\n".join(parts).strip()


def read_pdf_file(path: Path) -> str:
    if not path.exists():
        return ""

    parts: List[str] = []

    with pymupdf.open(path) as doc:
        for page in doc:
            text = page.get_text("text", sort=True)
            if text:
                parts.append(text.strip())

    return "\n".join(parts).strip()


def extract_docx_headings(path: Path) -> List[str]:
    if not path.exists():
        return []

    doc = Document(path)
    headings: List[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = ""
        try:
            style_name = paragraph.style.name.lower() if paragraph.style and paragraph.style.name else ""
        except Exception:
            style_name = ""

        if "heading" in style_name or "заголов" in style_name:
            headings.append(text)

    unique_headings: List[str] = []
    seen = set()
    for heading in headings:
        if heading not in seen:
            unique_headings.append(heading)
            seen.add(heading)

    return unique_headings


def looks_binary(path: Path) -> bool:
    try:
        chunk = path.read_bytes()[:4096]
    except Exception:
        return True

    if not chunk:
        return False

    if b"\x00" in chunk:
        return True

    text_chars = bytearray({7, 8, 9, 10, 12, 13, 27} | set(range(32, 127)))
    nontext = sum(byte not in text_chars for byte in chunk)
    return nontext / max(len(chunk), 1) > 0.30


def unique_paths(paths: List[Path]) -> List[Path]:
    result: List[Path] = []
    seen = set()

    for path in paths:
        try:
            key = str(path.resolve())
        except Exception:
            key = str(path)

        if key not in seen and path.exists():
            seen.add(key)
            result.append(path)

    return result


def collect_files_from_dir(dir_path: Path) -> List[Path]:
    if not dir_path.exists() or not dir_path.is_dir():
        return []
    return sorted([p for p in dir_path.rglob("*") if p.is_file()])


def collect_assignment_files() -> List[Path]:
    files = collect_files_from_dir(INPUT_DIR / "assignment")

    if not files and INPUT_DIR.exists():
        for file in sorted(INPUT_DIR.iterdir()):
            if (
                file.is_file()
                and file.name not in SPECIAL_ROOT_FILES
                and file.suffix.lower() in ASSIGNMENT_EXTENSIONS
            ):
                files.append(file)

    return unique_paths(files)


def collect_code_files() -> List[Path]:
    files = collect_files_from_dir(INPUT_DIR / "code")

    legacy_code = INPUT_DIR / "code.py"
    if legacy_code.exists():
        files.append(legacy_code)

    if not files and INPUT_DIR.exists():
        for file in sorted(INPUT_DIR.iterdir()):
            if (
                file.is_file()
                and file.name not in SPECIAL_ROOT_FILES
                and file.suffix.lower() in CODE_EXTENSIONS
            ):
                files.append(file)

    return unique_paths(files)


def collect_reference_files() -> List[Path]:
    files = collect_files_from_dir(INPUT_DIR / "reference")

    if not files and INPUT_DIR.exists():
        for file in sorted(INPUT_DIR.iterdir()):
            if not file.is_file():
                continue
            if file.name in SPECIAL_ROOT_FILES:
                continue
            if file.suffix.lower() in ASSIGNMENT_EXTENSIONS:
                continue
            if file.suffix.lower() in CODE_EXTENSIONS:
                continue
            if file.suffix.lower() in IMAGE_EXTENSIONS:
                continue
            files.append(file)

    return unique_paths(files)


def collect_image_files() -> List[Path]:
    files = collect_files_from_dir(INPUT_DIR / "images")

    if not files and INPUT_DIR.exists():
        for file in sorted(INPUT_DIR.iterdir()):
            if file.is_file() and file.suffix.lower() in IMAGE_EXTENSIONS:
                files.append(file)

    return unique_paths(files)


def read_assignment() -> Dict[str, object]:
    files = collect_assignment_files()

    if not files:
        return {
            "source_name": "",
            "text": "",
            "headings": [],
        }

    parts: List[str] = []
    headings: List[str] = []
    source_names: List[str] = []

    for file_path in files:
        suffix = file_path.suffix.lower()
        text = ""

        if suffix == ".txt":
            text = read_text_file_safe(file_path)
        elif suffix == ".docx":
            text = read_docx_file(file_path)
            headings.extend(extract_docx_headings(file_path))
        elif suffix == ".pdf":
            text = read_pdf_file(file_path)

        if text.strip():
            source_names.append(file_path.name)
            parts.append(f"[ФАЙЛ ЗАДАНИЯ: {file_path.name}]\n{text}")

    unique_headings: List[str] = []
    seen = set()
    for heading in headings:
        if heading not in seen:
            unique_headings.append(heading)
            seen.add(heading)

    return {
        "source_name": ", ".join(source_names),
        "text": "\n\n".join(parts).strip(),
        "headings": unique_headings,
    }


def read_code_entries() -> List[Dict[str, str]]:
    entries: List[Dict[str, str]] = []

    for file_path in collect_code_files():
        text = read_text_file_safe(file_path)
        if text.strip():
            entries.append(
                {
                    "name": file_path.name,
                    "text": text,
                }
            )

    return entries


def read_reference_entries() -> List[Dict[str, str]]:
    entries: List[Dict[str, str]] = []

    for file_path in collect_reference_files():
        if looks_binary(file_path):
            continue

        text = read_text_file_safe(file_path)
        if text.strip():
            entries.append(
                {
                    "name": file_path.name,
                    "text": text,
                }
            )

    return entries


def find_signature_image() -> Optional[Path]:
    keywords = ["sign", "signature", "podpis", "подпис"]

    for file in collect_image_files():
        lower_name = file.stem.lower()
        if any(keyword in lower_name for keyword in keywords):
            return file

    return None


def collect_non_signature_images() -> List[Path]:
    signature_image = find_signature_image()
    result: List[Path] = []

    for file in collect_image_files():
        if signature_image:
            try:
                if file.resolve() == signature_image.resolve():
                    continue
            except Exception:
                pass
        result.append(file)

    return result


# -----------------------------
# IMAGE ANALYSIS / SCALING
# -----------------------------
def normalize_dpi_value(value: Any, default: float = 96.0) -> float:
    try:
        value = float(value)
        if 10 <= value <= 1200:
            return value
    except Exception:
        pass
    return default


def get_image_natural_size_inches(path: Path, default_dpi: float = 96.0) -> Tuple[float, float]:
    try:
        with Image.open(path) as img:
            width_px, height_px = img.size
            dpi = img.info.get("dpi", (default_dpi, default_dpi))
            xdpi = normalize_dpi_value(dpi[0] if isinstance(dpi, tuple) else dpi, default_dpi)
            ydpi = normalize_dpi_value(dpi[1] if isinstance(dpi, tuple) else dpi, default_dpi)

            width_in = max(width_px / xdpi, 0.1)
            height_in = max(height_px / ydpi, 0.1)
            return width_in, height_in
    except Exception:
        return 4.0, 3.0


def fit_size_into_box(
    width_in: float,
    height_in: float,
    max_width_in: float,
    max_height_in: float,
    allow_upscale: bool = False,
) -> Tuple[float, float]:
    if width_in <= 0 or height_in <= 0:
        return max_width_in, min(max_height_in, max_width_in * 0.75)

    scale = min(max_width_in / width_in, max_height_in / height_in)
    if not allow_upscale:
        scale = min(scale, 1.0)

    return width_in * scale, height_in * scale


def get_docx_image_size(path: Path, max_width_in: float = 6.0, max_height_in: float = 7.0) -> Tuple[Any, Any]:
    width_in, height_in = get_image_natural_size_inches(path)
    fitted_w, fitted_h = fit_size_into_box(width_in, height_in, max_width_in, max_height_in, allow_upscale=False)
    return Inches(fitted_w), Inches(fitted_h)


def get_pdf_image_size(path: Path, max_width_cm: float = 15.0, max_height_cm: float = 18.0) -> Tuple[float, float]:
    width_in, height_in = get_image_natural_size_inches(path)
    width_cm = width_in * 2.54
    height_cm = height_in * 2.54
    fitted_w, fitted_h = fit_size_into_box(width_cm, height_cm, max_width_cm, max_height_cm, allow_upscale=False)
    return fitted_w * cm, fitted_h * cm


def get_gemini_image_part(path: Path) -> Optional[types.Part]:
    try:
        with Image.open(path) as img:
            img.load()

            max_dim = MAX_IMAGE_DIMENSION_FOR_ANALYSIS
            if max(img.size) > max_dim:
                img.thumbnail((max_dim, max_dim))

            suffix = path.suffix.lower()
            buf = io.BytesIO()

            if suffix in {".jpg", ".jpeg"}:
                if img.mode not in ("RGB", "L"):
                    img = img.convert("RGB")
                img.save(buf, format="JPEG", quality=85, optimize=True)
                mime = "image/jpeg"
            else:
                if img.mode not in ("RGB", "RGBA", "L"):
                    img = img.convert("RGBA")
                img.save(buf, format="PNG", optimize=True)
                mime = "image/png"

            return types.Part.from_bytes(
                data=buf.getvalue(),
                mime_type=mime,
            )
    except (UnidentifiedImageError, OSError, ValueError):
        return None


def iter_models(primary_model: str) -> List[str]:
    models = [primary_model] + [m for m in GEMINI_MODELS_FALLBACK if m != primary_model]
    seen = set()
    result = []
    for model in models:
        if model not in seen:
            seen.add(model)
            result.append(model)
    return result


def analyze_images_with_gemini(
    api_key: str,
    model_name: str,
    image_paths: List[Path],
) -> Dict[str, Dict[str, str]]:
    if not image_paths:
        return {}

    selected_images = image_paths[:MAX_IMAGES_FOR_ANALYSIS]
    parts_for_prompt = []
    actual_image_names: List[str] = []

    for idx, image_path in enumerate(selected_images, start=1):
        part = get_gemini_image_part(image_path)
        if part is None:
            continue

        actual_image_names.append(image_path.name)
        parts_for_prompt.append(f"Изображение {idx}. Имя файла: {image_path.name}")
        parts_for_prompt.append(part)

    if not parts_for_prompt:
        return {}

    instruction = """
Ты анализируешь изображения, которые пользователь приложил к лабораторной работе.

Для каждого изображения верни:
- filename
- summary: кратко опиши, что на изображении
- visible_text: видимый текст, если он есть
- relevance: может ли это быть полезно для отчёта и как именно

Верни СТРОГО валидный JSON вида:
{
  "images": [
    {
      "filename": "имя_файла",
      "summary": "краткое описание",
      "visible_text": "видимый текст или пусто",
      "relevance": "краткая оценка полезности"
    }
  ]
}

Правила:
- Не выдумывай то, чего не видно.
- Если изображение похоже на случайное фото и не относится к теме, так и скажи.
- Если это скриншот интерфейса, график, таблица, схема или формула, опиши это явно.
- Только JSON, без пояснений.
""".strip()

    client = genai.Client(api_key=api_key)
    last_error: Optional[Exception] = None

    for model in iter_models(model_name):
        for attempt in range(2):
            try:
                response = client.models.generate_content(
                    model=model,
                    contents=[instruction] + parts_for_prompt,
                    config=types.GenerateContentConfig(
                        temperature=0.2,
                        response_mime_type="application/json",
                    ),
                )

                text = (response.text or "").strip()
                data = json.loads(text)

                images = data.get("images", [])
                if not isinstance(images, list):
                    raise RuntimeError("Поле 'images' не является списком.")

                result: Dict[str, Dict[str, str]] = {}
                for item in images:
                    if not isinstance(item, dict):
                        continue

                    filename = str(item.get("filename", "")).strip()
                    if not filename:
                        continue

                    result[filename] = {
                        "summary": str(item.get("summary", "")).strip(),
                        "visible_text": str(item.get("visible_text", "")).strip(),
                        "relevance": str(item.get("relevance", "")).strip(),
                    }

                for filename in actual_image_names:
                    result.setdefault(
                        filename,
                        {"summary": "", "visible_text": "", "relevance": ""},
                    )

                return result

            except Exception as e:
                last_error = e
                if attempt == 0:
                    time.sleep(1.0)
                else:
                    break

    # Если анализ упал, просто молча возвращаем пустое — генерация лабы не должна из-за этого умирать.
    return {}


def format_image_analysis_for_prompt(image_analysis: Dict[str, Dict[str, str]]) -> str:
    if not image_analysis:
        return ""

    parts: List[str] = []

    for filename, info in image_analysis.items():
        summary = info.get("summary", "").strip()
        visible_text = info.get("visible_text", "").strip()
        relevance = info.get("relevance", "").strip()

        chunk = [f"[ИЗОБРАЖЕНИЕ: {filename}]"]
        if summary:
            chunk.append(f"Описание: {summary}")
        if visible_text:
            chunk.append(f"Видимый текст: {visible_text}")
        if relevance:
            chunk.append(f"Полезность: {relevance}")

        parts.append("\n".join(chunk))

    return "\n\n".join(parts).strip()


def build_image_caption(image_path: Path, idx: int, image_analysis: Dict[str, Dict[str, str]]) -> str:
    info = image_analysis.get(image_path.name, {})
    summary = (info.get("summary") or "").strip()
    visible_text = (info.get("visible_text") or "").strip()

    caption_text = ""
    if summary:
        caption_text = summary
    elif visible_text:
        caption_text = visible_text
    else:
        caption_text = image_path.stem

    caption_text = re.sub(r"\s+", " ", caption_text).strip()
    if len(caption_text) > 120:
        caption_text = caption_text[:117].rstrip() + "..."

    return f"Рисунок {idx} — {caption_text}"


# -----------------------------
# PROMPT / GEMINI
# -----------------------------
def format_entries_for_prompt(
    entries: List[Dict[str, str]],
    label: str,
    per_file_limit: int,
    total_limit: int,
) -> str:
    parts: List[str] = []
    total = 0

    for entry in entries:
        name = entry["name"]
        text = (entry["text"] or "").strip()
        if not text:
            continue

        clipped = text[:per_file_limit]
        chunk = f"[{label}: {name}]\n{clipped}"

        if total + len(chunk) > total_limit:
            remaining = total_limit - total
            if remaining <= 0:
                break

            head = f"[{label}: {name}]\n"
            allowed_text = max(0, remaining - len(head) - len("\n[ОБРЕЗАНО]"))
            if allowed_text > 0:
                chunk = f"{head}{clipped[:allowed_text]}\n[ОБРЕЗАНО]"
                parts.append(chunk)
            break

        parts.append(chunk)
        total += len(chunk)

    return "\n\n".join(parts).strip()


def build_prompt(
    assignment_text: str,
    detected_headings: List[str],
    notes: str,
    code_entries: List[Dict[str, str]],
    reference_entries: List[Dict[str, str]],
    image_analysis: Dict[str, Dict[str, str]],
    meta: Dict[str, Any],
) -> str:
    assignment_text = assignment_text[:30000]
    notes = notes[:8000]

    code_text = format_entries_for_prompt(
        code_entries,
        label="КОДОВЫЙ ФАЙЛ",
        per_file_limit=6000,
        total_limit=18000,
    )
    reference_text = format_entries_for_prompt(
        reference_entries,
        label="ДОПОЛНИТЕЛЬНЫЙ РЕФЕРЕНС",
        per_file_limit=5000,
        total_limit=15000,
    )
    image_text = format_image_analysis_for_prompt(image_analysis)

    detected_headings_json = json.dumps(detected_headings, ensure_ascii=False)
    subject = str(meta.get("subject", "")).strip()

    if not code_text.strip():
        code_instruction = """
Если код программы не предоставлен, но по смыслу лабораторной он действительно нужен, сгенерируй короткий и понятный студенческий пример.
Не превращай код в огромный мини-проект.
Если вставляешь код, оформи его как отдельную секцию с заголовком "Листинг программы" или "Код программы".
"""
    else:
        code_instruction = """
Если предоставлены кодовые файлы, используй их как основной источник для разделов про реализацию, алгоритм, листинг и пояснение работы программы.
Не нужно переписывать в отчёт весь код целиком, если он очень большой. Можно использовать ключевые фрагменты и кратко объяснять назначение файлов.
"""

    image_instruction = """
Если есть анализ изображений, используй его только там, где это реально уместно:
- для скриншотов интерфейса,
- схем,
- графиков,
- формул,
- таблиц,
- результатов программы.

Если изображение явно не относится к теме лабораторной, не встраивай его в основной текст насильно.
Но подписи к рисункам и приложение можно делать на основе анализа изображений.
"""

    return f"""
Ты помогаешь подготовить отчёт по лабораторной работе для Word-документа и PDF-документа.

Контекст:
- Дисциплина: {subject if subject else "не указана"}
- Это именно отчёт по лабораторной работе, а не обзорная статья и не реферат.

Главные правила:
- Если в методичке явно заданы названия разделов, используй их.
- Если найдены заголовки документа, ориентируйся на них в первую очередь.
- Построй структуру так, чтобы отчёт был похож на обычную сдаваемую студенческую лабораторную работу.
- Пиши по делу.
- Не используй канцелярские фразы вроде "в рамках данной работы", "следует отметить", "данная работа посвящена".
- Не делай текст слишком энциклопедическим.
- Если для демонстрации не хватает примеров, входных данных или кратких пояснений, можешь привести разумный учебный пример.
- Не пиши "данных недостаточно", если можно нормально достроить черновик по теме.
- Не добавляй лишние разделы.
- Не делай текст ультра-кратким: разделы должны быть достаточно содержательными.
- Если в тексте встречаются таблицы в markdown-формате через |, сохраняй эти данные как таблицы, не превращай их в кашу.
- Дополнительные референсы используй как вспомогательный материал: для терминов, структуры, пояснений, исходных данных и контекста.

Форматирование:
- Не используй markdown-заголовки через #.
- Не используй тройные кавычки ``` .
- Можно использовать только inline-выделение:
  **жирный** — для важных терминов
  *курсив* — для отдельных названий
  `код` — для коротких формул, векторов, параметров и фрагментов кода

Списки:
- Если перечисляешь несколько пунктов, каждый пункт пиши с новой строки.
- Не склеивай нумерованные списки в один абзац.
- Для перечислений используй формат:
  1. ...
  2. ...
  3. ...

Структура:
- Верни иерархию разделов.
- У раздела могут быть подразделы.
- Если подразделы уместны, обязательно используй их.
- Особенно если в работе есть блоки вроде "Базы данных", "Калькулятор CVSS", "Отличия версий", "Вывод" и т.п.

{code_instruction}

{image_instruction}

Найденные заголовки:
{detected_headings_json}

[ТЕКСТ МЕТОДИЧКИ / ЗАДАНИЯ]
{assignment_text if assignment_text.strip() else "Не предоставлено"}

[ЗАМЕТКИ]
{notes if notes.strip() else "Нет"}

[КОДОВЫЕ ФАЙЛЫ]
{code_text if code_text.strip() else "Не предоставлены"}

[ДОПОЛНИТЕЛЬНЫЕ РЕФЕРЕНСЫ]
{reference_text if reference_text.strip() else "Нет"}

[АНАЛИЗ ИЗОБРАЖЕНИЙ]
{image_text if image_text.strip() else "Нет"}

Верни СТРОГО валидный JSON формата:

{{
  "report_title": "название лабораторной работы",
  "sections": [
    {{
      "heading": "название раздела",
      "content": "текст раздела",
      "subsections": [
        {{
          "heading": "название подраздела",
          "content": "текст подраздела"
        }}
      ]
    }}
  ]
}}

Требования:
- Только JSON
- Без комментариев
- Без пояснений вне JSON
- Поле "subsections" можно делать пустым списком []
""".strip()


def generate_report_json_gemini(api_key: str, model_name: str, prompt: str) -> Dict[str, Any]:
    client = genai.Client(api_key=api_key)

    last_error: Optional[Exception] = None

    for model in iter_models(model_name):
        for attempt in range(3):
            try:
                response = client.models.generate_content(
                    model=model,
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        temperature=0.35,
                        response_mime_type="application/json",
                    ),
                )

                text = (response.text or "").strip()

                try:
                    data = json.loads(text)
                except json.JSONDecodeError as e:
                    raise RuntimeError(
                        f"Модель {model} вернула невалидный JSON:\n{text}"
                    ) from e

                if not isinstance(data, dict):
                    raise RuntimeError(f"Модель {model} вернула не JSON-объект.")

                if "sections" not in data or not isinstance(data["sections"], list):
                    raise RuntimeError(
                        f"Модель {model} вернула JSON без нормального поля 'sections'."
                    )

                normalize_sections(data["sections"])
                return data

            except Exception as e:
                last_error = e
                if attempt < 2:
                    time.sleep(1.5 * (attempt + 1))
                else:
                    break

    raise RuntimeError(f"Все Gemini-модели упали. Последняя ошибка:\n{last_error}")


def normalize_sections(sections: List[Dict[str, Any]]) -> None:
    for section in sections:
        if "heading" not in section:
            section["heading"] = ""
        if "content" not in section:
            section["content"] = ""
        if "subsections" not in section or not isinstance(section["subsections"], list):
            section["subsections"] = []

        for subsection in section["subsections"]:
            if "heading" not in subsection:
                subsection["heading"] = ""
            if "content" not in subsection:
                subsection["content"] = ""


# -----------------------------
# CLEANUP / PARSING
# -----------------------------
def cleanup_heading(text: str) -> str:
    text = text.replace("**", "").replace("*", "").replace("`", "")
    text = re.sub(r"^\s*#{1,6}\s*", "", text, flags=re.MULTILINE)
    return text.strip()


def cleanup_code_fences(text: str) -> str:
    if not text:
        return text

    text = text.strip()

    if text.startswith("```"):
        lines = text.splitlines()
        if lines:
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        return "\n".join(lines).strip()

    return text


def is_pipe_table_line(line: str) -> bool:
    stripped = line.strip()
    return stripped.startswith("|") and stripped.endswith("|") and stripped.count("|") >= 2


def split_into_blocks(text: str) -> List[str]:
    if not text.strip():
        return []

    lines = [line.rstrip() for line in text.splitlines()]

    result: List[str] = []
    paragraph_buffer: List[str] = []
    table_buffer: List[str] = []

    list_item_pattern = re.compile(r"^\s*(\d+[\.\)]|[-•])\s+")
    heading_like_pattern = re.compile(r"^\s*\d+(\.\d+)*\.\s+")

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            result.append(" ".join(x.strip() for x in paragraph_buffer if x.strip()).strip())
            paragraph_buffer = []

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            result.append("\n".join(table_buffer).strip())
            table_buffer = []

    for line in lines:
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            flush_table()
            continue

        if is_pipe_table_line(stripped):
            flush_paragraph()
            table_buffer.append(stripped)
            continue

        if table_buffer:
            flush_table()

        if list_item_pattern.match(stripped) or heading_like_pattern.match(stripped):
            flush_paragraph()
            result.append(stripped)
            continue

        paragraph_buffer.append(stripped)

    flush_paragraph()
    flush_table()
    return result


def split_pipe_row(line: str) -> List[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def parse_pipe_table(block: str) -> Optional[Dict[str, Any]]:
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    if len(lines) < 2:
        return None

    if not is_pipe_table_line(lines[0]):
        return None

    if not PIPE_TABLE_SEPARATOR_RE.match(lines[1]):
        return None

    headers = split_pipe_row(lines[0])
    if not headers:
        return None

    width = len(headers)
    rows: List[List[str]] = []

    for line in lines[2:]:
        if not is_pipe_table_line(line):
            return None

        cells = split_pipe_row(line)

        if len(cells) < width:
            cells += [""] * (width - len(cells))
        elif len(cells) > width:
            cells = cells[:width]

        rows.append(cells)

    return {
        "headers": headers,
        "rows": rows,
    }


def parse_inline_list_items(text: str) -> List[str]:
    raw = text.strip()
    normalized = re.sub(r"[ \t]+", " ", raw)

    matches = list(re.finditer(r"(?=(?:^|\s)(\d+)\.\s+)", normalized))
    if len(matches) >= 2:
        items = []
        for i, match in enumerate(matches):
            start = match.start()
            if start > 0 and normalized[start].isspace():
                start += 1
            end = matches[i + 1].start() if i + 1 < len(matches) else len(normalized)
            chunk = normalized[start:end].strip()
            if chunk:
                items.append(chunk)
        return items

    metric_heads = [
        "Базовые метрики",
        "Временные метрики",
        "Контекстные метрики",
        "Base Metrics",
        "Temporal Metrics",
        "Environmental Metrics",
    ]

    positions = []
    for head in metric_heads:
        for match in re.finditer(re.escape(head), normalized):
            positions.append((match.start(), head))

    positions.sort(key=lambda x: x[0])

    if len(positions) >= 2:
        items = []
        for i, (start, _) in enumerate(positions):
            end = positions[i + 1][0] if i + 1 < len(positions) else len(normalized)
            chunk = normalized[start:end].strip(" :;")
            if chunk:
                items.append(f"• {chunk}")
        return items

    return []


def extract_code_from_sections(report_data: Dict[str, Any]) -> Optional[str]:
    code_like_titles = {
        "листинг программы",
        "код программы",
        "программа",
        "листинг",
        "исходный код",
    }

    sections = report_data.get("sections", [])
    cleaned_sections = []
    extracted_code = None

    for section in sections:
        heading = str(section.get("heading", "")).strip().lower()
        content = str(section.get("content", "")).strip()

        if heading in code_like_titles and content:
            extracted_code = cleanup_code_fences(content)
            continue

        sub_cleaned = []
        for subsection in section.get("subsections", []):
            sub_heading = str(subsection.get("heading", "")).strip().lower()
            sub_content = str(subsection.get("content", "")).strip()

            if sub_heading in code_like_titles and sub_content and not extracted_code:
                extracted_code = cleanup_code_fences(sub_content)
                continue

            sub_cleaned.append(subsection)

        section["subsections"] = sub_cleaned
        cleaned_sections.append(section)

    report_data["sections"] = cleaned_sections
    return extracted_code


# -----------------------------
# DOCX FORMATTING
# -----------------------------
def set_run_font(run, font_name: str = "Times New Roman", font_size: int = 14) -> None:
    run.font.name = font_name
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(font_size)


def apply_gost_style(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(14)

    normal_pf = normal_style.paragraph_format
    normal_pf.first_line_indent = Cm(1.25)
    normal_pf.line_spacing = 1.5
    normal_pf.space_before = Pt(0)
    normal_pf.space_after = Pt(0)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(14)
            style.font.bold = True

            pf = style.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.left_indent = Cm(0)
            pf.line_spacing = 1.5
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)


def add_formatted_text(paragraph: Paragraph, text: str) -> None:
    parts = INLINE_PATTERN.split(text)

    for part in parts:
        if not part:
            continue

        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            content = part[2:-2]
            run = paragraph.add_run(content)
            run.bold = True
            set_run_font(run, "Times New Roman", 14)

        elif (
            part.startswith("*")
            and part.endswith("*")
            and len(part) >= 2
            and not (part.startswith("**") and part.endswith("**"))
        ):
            content = part[1:-1]
            run = paragraph.add_run(content)
            run.italic = True
            set_run_font(run, "Times New Roman", 14)

        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            content = part[1:-1]
            run = paragraph.add_run(content)
            set_run_font(run, "Courier New", 11)

        else:
            run = paragraph.add_run(part)
            set_run_font(run, "Times New Roman", 14)


def add_list_item_paragraph(doc: Document, text: str) -> None:
    match = LIST_ITEM_RE.match(text.strip())
    if match:
        marker = match.group(1)
        content = match.group(2).strip()
    else:
        marker = "•"
        content = text.strip()
        if content.startswith("•"):
            content = content[1:].strip()

    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5

    run = p.add_run(f"{marker} ")
    set_run_font(run, "Times New Roman", 14)

    add_formatted_text(p, content)


def fill_table_cell(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if bold:
        run = p.add_run(text)
        run.bold = True
        set_run_font(run, "Times New Roman", 12)
    else:
        add_formatted_text(p, text)


def add_docx_table(doc: Document, headers: List[str], rows: List[List[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        fill_table_cell(header_cells[i], header, bold=True)

    for row in rows:
        row_cells = table.add_row().cells
        for i, value in enumerate(row):
            fill_table_cell(row_cells[i], value, bold=False)

    doc.add_paragraph()


def render_content_block(doc: Document, block: str) -> None:
    block = block.strip()
    if not block:
        return

    table_data = parse_pipe_table(block)
    if table_data:
        add_docx_table(doc, table_data["headers"], table_data["rows"])
        return

    inline_items = parse_inline_list_items(block)
    if inline_items:
        for item in inline_items:
            add_list_item_paragraph(doc, item)
        return

    if LIST_ITEM_RE.match(block):
        add_list_item_paragraph(doc, block)
        return

    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    add_formatted_text(p, block)


def add_paragraph_text(doc: Document, text: str) -> None:
    for block in split_into_blocks(text):
        render_content_block(doc, block)


def add_heading_paragraph(doc: Document, text: str, level: int = 1) -> None:
    safe_level = min(max(level, 1), 3)
    p = doc.add_paragraph(style=f"Heading {safe_level}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5

    run = p.add_run(text)
    run.bold = True
    set_run_font(run, "Times New Roman", 14)


def add_manual_table_of_contents(doc: Document, sections: List[Dict[str, Any]]) -> None:
    add_heading_paragraph(doc, "Содержание", level=1)

    for i, section in enumerate(sections, start=1):
        heading = cleanup_heading(str(section.get("heading", "")).strip())
        if heading:
            p = doc.add_paragraph()
            p.style = doc.styles["Normal"]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.5
            add_formatted_text(p, f"{i}. {heading}")

        for j, subsection in enumerate(section.get("subsections", []), start=1):
            sub_heading = cleanup_heading(str(subsection.get("heading", "")).strip())
            if sub_heading:
                p = doc.add_paragraph()
                p.style = doc.styles["Normal"]
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.line_spacing = 1.5
                add_formatted_text(p, f"{i}.{j}. {sub_heading}")


def add_title_page(doc: Document, meta: Dict[str, Any], report_title: str) -> None:
    university = str(meta.get("university", "")).strip()
    faculty = str(meta.get("faculty", "")).strip()
    subject = str(meta.get("subject", "")).strip()
    student_name = str(meta.get("student_name", "")).strip()
    group = str(meta.get("group", "")).strip()
    teacher = str(meta.get("teacher", "")).strip()
    city = str(meta.get("city", "")).strip()
    year = str(meta.get("year", "")).strip()

    signature_image = find_signature_image()

    lab_header = "ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ"
    lab_topic = ""

    cleaned_title = report_title.strip()
    match = re.match(
        r"^\s*Лабораторная\s+работа\s*№\s*(\d+)\s*[:\-]?\s*(.*)$",
        cleaned_title,
        flags=re.IGNORECASE,
    )
    if match:
        number = match.group(1).strip()
        tail = match.group(2).strip()
        lab_header = f"ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № {number}"
        lab_topic = tail
    else:
        lab_topic = cleaned_title

    if university:
        uni_lines = [line.strip() for line in university.split("\n") if line.strip()]
        for idx, line in enumerate(uni_lines):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15

            run = p.add_run(line)
            run.bold = True
            set_run_font(run, "Times New Roman", 8 if idx == 0 else 12)

    if faculty:
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.15

        run = p.add_run(faculty)
        run.bold = True
        set_run_font(run, "Times New Roman", 12)

    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run("Дисциплина:")
    run.bold = True
    set_run_font(run, "Times New Roman", 12)

    if subject:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(f"«{subject}»")
        set_run_font(run, "Times New Roman", 12)

    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(lab_header)
    run.bold = True
    set_run_font(run, "Times New Roman", 12)

    if lab_topic:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(f"«{lab_topic}»")
        set_run_font(run, "Times New Roman", 12)

    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run("Выполнил:")
    run.bold = True
    set_run_font(run, "Times New Roman", 12)

    if group:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(f"Студент гр. {group}")
        set_run_font(run, "Times New Roman", 12)

    if student_name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(student_name)
        set_run_font(run, "Times New Roman", 12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run("Подпись:")
    set_run_font(run, "Times New Roman", 12)

    if signature_image:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run()
        run.add_picture(str(signature_image), height=Cm(1.5))

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run("Проверил:")
    run.bold = True
    set_run_font(run, "Times New Roman", 12)

    if teacher:
        teacher_lines = [line.strip() for line in teacher.split("\n") if line.strip()]
        for line in teacher_lines:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line)
            set_run_font(run, "Times New Roman", 12)

    for _ in range(7):
        doc.add_paragraph()

    footer = ""
    if city and year:
        footer = f"{city}\n{year}г."
    elif city:
        footer = city
    elif year:
        footer = f"{year}г."

    if footer:
        for line in footer.split("\n"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line)
            set_run_font(run, "Times New Roman", 12)

    doc.add_page_break()


def render_sections(doc: Document, sections: List[Dict[str, Any]]) -> None:
    for i, section in enumerate(sections, start=1):
        heading = cleanup_heading(str(section.get("heading", "")).strip())
        content = str(section.get("content", "")).strip()
        subsections = section.get("subsections", [])

        if heading:
            add_heading_paragraph(doc, f"{i}. {heading}", level=1)

        if content:
            add_paragraph_text(doc, content)

        for j, subsection in enumerate(subsections, start=1):
            sub_heading = cleanup_heading(str(subsection.get("heading", "")).strip())
            sub_content = str(subsection.get("content", "")).strip()

            if sub_heading:
                add_heading_paragraph(doc, f"{i}.{j}. {sub_heading}", level=2)

            if sub_content:
                add_paragraph_text(doc, sub_content)


def add_code_section(doc: Document, code_text: str) -> None:
    if not code_text.strip():
        return

    doc.add_page_break()
    add_heading_paragraph(doc, "Приложение А. Листинг программы", level=1)

    code_text = cleanup_code_fences(code_text)

    for line in code_text.splitlines():
        p = doc.add_paragraph()
        p.style = doc.styles["Normal"]
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0

        run = p.add_run(line if line else " ")
        set_run_font(run, "Courier New", 10)


def add_images_from_input_docx(doc: Document, image_analysis: Dict[str, Dict[str, str]]) -> None:
    images = collect_non_signature_images()

    if not images:
        return

    doc.add_page_break()
    add_heading_paragraph(doc, "Приложение Б. Иллюстрации", level=1)

    for idx, image_path in enumerate(images, start=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)

        width, height = get_docx_image_size(image_path, max_width_in=6.0, max_height_in=7.2)

        run = p.add_run()
        run.add_picture(str(image_path), width=width, height=height)

        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.first_line_indent = Cm(0)

        caption = build_image_caption(image_path, idx, image_analysis)
        run = cap.add_run(caption)
        set_run_font(run, "Times New Roman", 12)


def create_docx(
    report_data: Dict[str, Any],
    code_text: str,
    meta: Dict[str, Any],
    output_path: Path,
    image_analysis: Dict[str, Dict[str, str]],
) -> None:
    doc = Document()
    apply_gost_style(doc)

    report_title = str(report_data.get("report_title", "Лабораторная работа")).strip()

    add_title_page(doc, meta, report_title)

    sections = report_data.get("sections", [])
    add_manual_table_of_contents(doc, sections)
    doc.add_page_break()

    render_sections(doc, sections)

    if code_text.strip():
        add_code_section(doc, code_text)

    add_images_from_input_docx(doc, image_analysis)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


# -----------------------------
# PDF HELPERS
# -----------------------------
def find_font_path(candidates: List[str]) -> Optional[Path]:
    search_dirs = [
        BASE_DIR / "fonts",
        BASE_DIR.parent / "fonts",
        Path("/usr/share/fonts/truetype/dejavu"),
        Path("/usr/share/fonts/dejavu"),
        Path("/usr/share/fonts/truetype/liberation2"),
        Path("/usr/share/fonts/truetype/liberation"),
        Path("/usr/share/fonts/TTF"),
        Path("/usr/local/share/fonts"),
        Path.home() / ".fonts",
        Path.home() / ".local/share/fonts",
    ]

    for directory in search_dirs:
        for name in candidates:
            path = directory / name
            if path.exists():
                return path

    return None


def register_pdf_fonts() -> Dict[str, str]:
    serif_regular_path = find_font_path(
        ["DejaVuSerif.ttf", "LiberationSerif-Regular.ttf", "FreeSerif.ttf"]
    )
    serif_bold_path = find_font_path(
        ["DejaVuSerif-Bold.ttf", "LiberationSerif-Bold.ttf", "FreeSerifBold.ttf"]
    )
    serif_italic_path = find_font_path(
        ["DejaVuSerif-Italic.ttf", "LiberationSerif-Italic.ttf", "FreeSerifItalic.ttf"]
    )
    mono_path = find_font_path(
        ["DejaVuSansMono.ttf", "LiberationMono-Regular.ttf", "FreeMono.ttf"]
    )

    fonts = {
        "regular": "Helvetica",
        "bold": "Helvetica-Bold",
        "italic": "Helvetica-Oblique",
        "mono": "Courier",
    }

    if serif_regular_path:
        pdfmetrics.registerFont(TTFont("LabSerif", str(serif_regular_path)))
        fonts["regular"] = "LabSerif"

    if serif_bold_path:
        pdfmetrics.registerFont(TTFont("LabSerifBold", str(serif_bold_path)))
        fonts["bold"] = "LabSerifBold"

    if serif_italic_path:
        pdfmetrics.registerFont(TTFont("LabSerifItalic", str(serif_italic_path)))
        fonts["italic"] = "LabSerifItalic"

    if mono_path:
        pdfmetrics.registerFont(TTFont("LabMono", str(mono_path)))
        fonts["mono"] = "LabMono"

    return fonts


def inline_to_reportlab_markup(text: str, fonts: Dict[str, str]) -> str:
    parts = INLINE_PATTERN.split(text)
    out: List[str] = []

    for part in parts:
        if not part:
            continue

        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            content = html.escape(part[2:-2], quote=False)
            out.append(f"<b>{content}</b>")

        elif (
            part.startswith("*")
            and part.endswith("*")
            and len(part) >= 2
            and not (part.startswith("**") and part.endswith("**"))
        ):
            content = html.escape(part[1:-1], quote=False)
            out.append(f"<i>{content}</i>")

        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            content = html.escape(part[1:-1], quote=False)
            out.append(f'<font name="{fonts["mono"]}">{content}</font>')

        else:
            out.append(html.escape(part, quote=False))

    return "".join(out)


def build_pdf_styles(fonts: Dict[str, str]) -> Dict[str, ParagraphStyle]:
    base_styles = getSampleStyleSheet()

    return {
        "small_center": ParagraphStyle(
            "small_center",
            parent=base_styles["Normal"],
            fontName=fonts["bold"],
            fontSize=8,
            leading=10,
            alignment=TA_CENTER,
            spaceAfter=4,
        ),
        "center": ParagraphStyle(
            "center",
            parent=base_styles["Normal"],
            fontName=fonts["regular"],
            fontSize=12,
            leading=15,
            alignment=TA_CENTER,
            spaceAfter=4,
        ),
        "center_bold": ParagraphStyle(
            "center_bold",
            parent=base_styles["Normal"],
            fontName=fonts["bold"],
            fontSize=12,
            leading=15,
            alignment=TA_CENTER,
            spaceAfter=4,
        ),
        "right": ParagraphStyle(
            "right",
            parent=base_styles["Normal"],
            fontName=fonts["regular"],
            fontSize=12,
            leading=15,
            alignment=TA_RIGHT,
            spaceAfter=4,
        ),
        "right_bold": ParagraphStyle(
            "right_bold",
            parent=base_styles["Normal"],
            fontName=fonts["bold"],
            fontSize=12,
            leading=15,
            alignment=TA_RIGHT,
            spaceAfter=4,
        ),
        "heading": ParagraphStyle(
            "heading",
            parent=base_styles["Normal"],
            fontName=fonts["bold"],
            fontSize=14,
            leading=18,
            alignment=TA_LEFT,
            spaceBefore=8,
            spaceAfter=6,
            firstLineIndent=0,
        ),
        "body": ParagraphStyle(
            "body",
            parent=base_styles["Normal"],
            fontName=fonts["regular"],
            fontSize=12,
            leading=18,
            alignment=TA_JUSTIFY,
            firstLineIndent=1.25 * cm,
            spaceAfter=4,
        ),
        "list_item": ParagraphStyle(
            "list_item",
            parent=base_styles["Normal"],
            fontName=fonts["regular"],
            fontSize=12,
            leading=18,
            alignment=TA_JUSTIFY,
            leftIndent=1.25 * cm,
            firstLineIndent=0,
            spaceAfter=2,
        ),
        "table_cell": ParagraphStyle(
            "table_cell",
            parent=base_styles["Normal"],
            fontName=fonts["regular"],
            fontSize=10,
            leading=12,
            alignment=TA_LEFT,
            spaceAfter=0,
            spaceBefore=0,
        ),
        "table_header": ParagraphStyle(
            "table_header",
            parent=base_styles["Normal"],
            fontName=fonts["bold"],
            fontSize=10,
            leading=12,
            alignment=TA_LEFT,
            spaceAfter=0,
            spaceBefore=0,
        ),
        "caption": ParagraphStyle(
            "caption",
            parent=base_styles["Normal"],
            fontName=fonts["regular"],
            fontSize=10,
            leading=12,
            alignment=TA_CENTER,
            spaceAfter=8,
        ),
        "mono": ParagraphStyle(
            "mono",
            parent=base_styles["Code"],
            fontName=fonts["mono"],
            fontSize=9,
            leading=11,
            leftIndent=0,
            rightIndent=0,
            spaceAfter=2,
        ),
    }


def pdf_add_title_page(
    story: List[Any],
    report_title: str,
    meta: Dict[str, Any],
    styles: Dict[str, ParagraphStyle],
) -> None:
    university = str(meta.get("university", "")).strip()
    faculty = str(meta.get("faculty", "")).strip()
    subject = str(meta.get("subject", "")).strip()
    student_name = str(meta.get("student_name", "")).strip()
    group = str(meta.get("group", "")).strip()
    teacher = str(meta.get("teacher", "")).strip()
    city = str(meta.get("city", "")).strip()
    year = str(meta.get("year", "")).strip()

    lab_header = "ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ"
    lab_topic = ""

    cleaned_title = report_title.strip()
    match = re.match(
        r"^\s*Лабораторная\s+работа\s*№\s*(\d+)\s*[:\-]?\s*(.*)$",
        cleaned_title,
        flags=re.IGNORECASE,
    )
    if match:
        number = match.group(1).strip()
        tail = match.group(2).strip()
        lab_header = f"ОТЧЕТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № {number}"
        lab_topic = tail
    else:
        lab_topic = cleaned_title

    if university:
        uni_lines = [line.strip() for line in university.split("\n") if line.strip()]
        for idx, line in enumerate(uni_lines):
            style = styles["small_center"] if idx == 0 else styles["center_bold"]
            story.append(RLParagraph(html.escape(line), style))

    if faculty:
        story.append(Spacer(1, 0.2 * cm))
        story.append(RLParagraph(html.escape(faculty), styles["center_bold"]))

    story.append(Spacer(1, 3.5 * cm))
    story.append(RLParagraph("Дисциплина:", styles["center_bold"]))
    if subject:
        story.append(RLParagraph(f"«{html.escape(subject)}»", styles["center"]))

    story.append(Spacer(1, 2.2 * cm))
    story.append(RLParagraph(html.escape(lab_header), styles["center_bold"]))
    if lab_topic:
        story.append(RLParagraph(f"«{html.escape(lab_topic)}»", styles["center"]))

    story.append(Spacer(1, 3 * cm))
    story.append(RLParagraph("Выполнил:", styles["right_bold"]))
    if group:
        story.append(RLParagraph(f"Студент гр. {html.escape(group)}", styles["right"]))
    if student_name:
        story.append(RLParagraph(html.escape(student_name), styles["right"]))

    signature_image = find_signature_image()
    if signature_image and signature_image.exists():
        try:
            img = RLImage(str(signature_image), width=4 * cm, height=1.2 * cm)
            img.hAlign = "RIGHT"
            story.append(RLParagraph("Подпись:", styles["right"]))
            story.append(img)
        except Exception:
            story.append(RLParagraph("Подпись: __________________", styles["right"]))
    else:
        story.append(RLParagraph("Подпись: __________________", styles["right"]))

    story.append(Spacer(1, 0.6 * cm))
    story.append(RLParagraph("Проверил:", styles["right_bold"]))
    if teacher:
        for line in teacher.split("\n"):
            line = line.strip()
            if line:
                story.append(RLParagraph(html.escape(line), styles["right"]))

    story.append(Spacer(1, 4 * cm))
    footer_parts = []
    if city:
        footer_parts.append(html.escape(city))
    if year:
        footer_parts.append(f"{html.escape(year)}г.")
    if footer_parts:
        story.append(RLParagraph("<br/>".join(footer_parts), styles["center"]))

    story.append(PageBreak())


def add_pdf_table(
    story: List[Any],
    headers: List[str],
    rows: List[List[str]],
    styles: Dict[str, ParagraphStyle],
    fonts: Dict[str, str],
) -> None:
    table_data: List[List[Any]] = []

    header_row = [
        RLParagraph(inline_to_reportlab_markup(header, fonts), styles["table_header"])
        for header in headers
    ]
    table_data.append(header_row)

    for row in rows:
        table_data.append(
            [
                RLParagraph(inline_to_reportlab_markup(cell, fonts), styles["table_cell"])
                for cell in row
            ]
        )

    table = RLTable(table_data, repeatRows=1)
    table.setStyle(
        RLTableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDEDED")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )

    story.append(table)
    story.append(Spacer(1, 0.2 * cm))


def pdf_add_content_block(
    story: List[Any],
    block: str,
    styles: Dict[str, ParagraphStyle],
    fonts: Dict[str, str],
) -> None:
    block = block.strip()
    if not block:
        return

    table_data = parse_pipe_table(block)
    if table_data:
        add_pdf_table(story, table_data["headers"], table_data["rows"], styles, fonts)
        return

    inline_items = parse_inline_list_items(block)
    if inline_items:
        for item in inline_items:
            pdf_add_content_block(story, item, styles, fonts)
        return

    match = LIST_ITEM_RE.match(block)
    if match:
        marker = html.escape(match.group(1), quote=False)
        content = inline_to_reportlab_markup(match.group(2).strip(), fonts)
        story.append(RLParagraph(f"{marker} {content}", styles["list_item"]))
        return

    story.append(RLParagraph(inline_to_reportlab_markup(block, fonts), styles["body"]))


def pdf_add_paragraph_text(
    story: List[Any],
    text: str,
    styles: Dict[str, ParagraphStyle],
    fonts: Dict[str, str],
) -> None:
    for block in split_into_blocks(text):
        pdf_add_content_block(story, block, styles, fonts)


def pdf_render_sections(
    story: List[Any],
    sections: List[Dict[str, Any]],
    styles: Dict[str, ParagraphStyle],
    fonts: Dict[str, str],
) -> None:
    for i, section in enumerate(sections, start=1):
        heading = cleanup_heading(str(section.get("heading", "")).strip())
        content = str(section.get("content", "")).strip()
        subsections = section.get("subsections", [])

        if heading:
            story.append(RLParagraph(html.escape(f"{i}. {heading}"), styles["heading"]))

        if content:
            pdf_add_paragraph_text(story, content, styles, fonts)

        for j, subsection in enumerate(subsections, start=1):
            sub_heading = cleanup_heading(str(subsection.get("heading", "")).strip())
            sub_content = str(subsection.get("content", "")).strip()

            if sub_heading:
                story.append(RLParagraph(html.escape(f"{i}.{j}. {sub_heading}"), styles["heading"]))

            if sub_content:
                pdf_add_paragraph_text(story, sub_content, styles, fonts)


def pdf_add_code_section(story: List[Any], code_text: str, styles: Dict[str, ParagraphStyle]) -> None:
    if not code_text.strip():
        return

    story.append(PageBreak())
    story.append(RLParagraph("Приложение А. Листинг программы", styles["heading"]))
    story.append(Preformatted(cleanup_code_fences(code_text), styles["mono"]))


def pdf_add_images_section(story: List[Any], styles: Dict[str, ParagraphStyle], image_analysis: Dict[str, Dict[str, str]]) -> None:
    images = collect_non_signature_images()

    if not images:
        return

    story.append(PageBreak())
    story.append(RLParagraph("Приложение Б. Иллюстрации", styles["heading"]))

    for idx, image_path in enumerate(images, start=1):
        try:
            width, height = get_pdf_image_size(image_path, max_width_cm=15.0, max_height_cm=18.0)
            img = RLImage(str(image_path), width=width, height=height)
            img.hAlign = "CENTER"
            story.append(img)
            story.append(
                RLParagraph(
                    html.escape(build_image_caption(image_path, idx, image_analysis)),
                    styles["caption"],
                )
            )
            story.append(Spacer(1, 0.3 * cm))
        except Exception:
            continue


def create_pdf(
    report_data: Dict[str, Any],
    code_text: str,
    meta: Dict[str, Any],
    output_path: Path,
    image_analysis: Dict[str, Dict[str, str]],
) -> None:
    fonts = register_pdf_fonts()
    styles = build_pdf_styles(fonts)

    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        leftMargin=3 * cm,
        rightMargin=1.5 * cm,
        title=str(report_data.get("report_title", "Лабораторная работа")).strip(),
    )

    story: List[Any] = []

    report_title = str(report_data.get("report_title", "Лабораторная работа")).strip()

    pdf_add_title_page(story, report_title, meta, styles)
    pdf_render_sections(story, report_data.get("sections", []), styles, fonts)

    if code_text.strip():
        pdf_add_code_section(story, code_text, styles)

    pdf_add_images_section(story, styles, image_analysis)

    doc.build(story)


# -----------------------------
# TELEGRAM OUTPUT
# -----------------------------
def split_long_text(text: str, limit: int) -> List[str]:
    clean_text = (text or "").strip()
    if not clean_text:
        return []

    paragraphs = [part.strip() for part in clean_text.split("\n") if part.strip()]
    if not paragraphs:
        paragraphs = [clean_text]

    pieces: List[str] = []
    current = ""

    for paragraph in paragraphs:
        candidate = paragraph if not current else f"{current}\n{paragraph}"
        if len(candidate) <= limit:
            current = candidate
            continue

        if current:
            pieces.append(current)
            current = ""

        if len(paragraph) <= limit:
            current = paragraph
            continue

        start = 0
        while start < len(paragraph):
            pieces.append(paragraph[start:start + limit])
            start += limit

    if current:
        pieces.append(current)

    return pieces


def report_to_telegram_chunks(report_data: Dict[str, Any], max_len: int = 3800) -> List[str]:
    def esc(text: str) -> str:
        return html.escape(text or "", quote=False)

    chunks: List[str] = []
    current = ""

    report_title = esc(str(report_data.get("report_title", "")).strip())
    if report_title:
        current = f"<b>{report_title}</b>\n\n"

    def flush_current() -> None:
        nonlocal current
        if current.strip():
            chunks.append(current.strip())
            current = ""

    def add_block(block: str) -> None:
        nonlocal current
        if len(block) > max_len:
            flush_current()
            chunks.append(block[:max_len].strip())
            rest = block[max_len:].strip()
            if rest:
                add_block(rest)
            return

        if len(current) + len(block) > max_len:
            flush_current()

        current += block

    sections = report_data.get("sections", [])

    for i, section in enumerate(sections, start=1):
        heading = esc(str(section.get("heading", "")).strip())
        content = str(section.get("content", "")).strip()

        if heading:
            add_block(f"<b>{i}. {heading}</b>\n")

        if content:
            content_parts = split_long_text(content, 3000)
            for part in content_parts:
                add_block(f"<blockquote expandable>{esc(part)}</blockquote>\n")

        for j, subsection in enumerate(section.get("subsections", []), start=1):
            sub_heading = esc(str(subsection.get("heading", "")).strip())
            sub_content = str(subsection.get("content", "")).strip()

            if sub_heading:
                add_block(f"<b>{i}.{j}. {sub_heading}</b>\n")

            if sub_content:
                sub_parts = split_long_text(sub_content, 3000)
                for part in sub_parts:
                    add_block(f"<blockquote expandable>{esc(part)}</blockquote>\n")

        add_block("\n")

    flush_current()
    return chunks


# -----------------------------
# MAIN
# -----------------------------
def generate_lab_report_from_dirs(input_dir: Path, output_dir: Path) -> Dict[str, Any]:
    load_dotenv()

    global INPUT_DIR, OUTPUT_DIR
    old_input, old_output = INPUT_DIR, OUTPUT_DIR

    INPUT_DIR = Path(input_dir)
    OUTPUT_DIR = Path(output_dir)

    try:
        gemini_api_key = os.getenv("GEMINI_API_KEY")
        gemini_model = os.getenv("GEMINI_MODEL", "gemini-2.5-flash-lite")

        if not gemini_api_key:
            raise RuntimeError("Не найден GEMINI_API_KEY в .env")

        meta = read_json_file(INPUT_DIR / "meta.json")

        assignment_data = read_assignment()
        assignment_text = str(assignment_data.get("text", "")).strip()
        detected_headings = assignment_data.get("headings", [])

        notes = read_text_file(INPUT_DIR / "notes.txt")
        code_entries = read_code_entries()
        reference_entries = read_reference_entries()
        image_analysis = analyze_images_with_gemini(
            api_key=gemini_api_key,
            model_name=gemini_model,
            image_paths=collect_non_signature_images(),
        )

        if not assignment_text and not notes and not code_entries and not reference_entries and not image_analysis:
            raise RuntimeError("Нет данных для генерации отчёта")

        prompt = build_prompt(
            assignment_text=assignment_text,
            detected_headings=detected_headings,
            notes=notes,
            code_entries=code_entries,
            reference_entries=reference_entries,
            image_analysis=image_analysis,
            meta=meta,
        )

        report_data = generate_report_json_gemini(
            api_key=gemini_api_key,
            model_name=gemini_model,
            prompt=prompt,
        )

        code_text = format_entries_for_prompt(
            code_entries,
            label="КОДОВЫЙ ФАЙЛ",
            per_file_limit=10000,
            total_limit=25000,
        )

        if not code_text.strip():
            extracted_code = extract_code_from_sections(report_data)
            if extracted_code:
                code_text = extracted_code

        output_docx = OUTPUT_DIR / "lab_report.docx"
        output_pdf = OUTPUT_DIR / "lab_report.pdf"

        create_docx(
            report_data=report_data,
            code_text=code_text,
            meta=meta,
            output_path=output_docx,
            image_analysis=image_analysis,
        )

        create_pdf(
            report_data=report_data,
            code_text=code_text,
            meta=meta,
            output_path=output_pdf,
            image_analysis=image_analysis,
        )

        telegram_chunks = report_to_telegram_chunks(report_data)

        return {
            "docx_path": output_docx,
            "pdf_path": output_pdf,
            "report_data": report_data,
            "telegram_chunks": telegram_chunks,
        }

    finally:
        INPUT_DIR = old_input
        OUTPUT_DIR = old_output


def main() -> None:
    result = generate_lab_report_from_dirs(INPUT_DIR, OUTPUT_DIR)
    print(f"Готово: {result['docx_path']}")
    print(f"PDF: {result['pdf_path']}")


if __name__ == "__main__":
    main()